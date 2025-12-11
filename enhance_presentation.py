"""
تحسين عرض بوربوينت راصد الموجود
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

# فتح العرض الموجود
prs = Presentation(r'C:\Users\Administrator\Desktop\Rased_Presentation.pptx')

# ألوان المشروع
PRIMARY_COLOR = RGBColor(16, 185, 129)
SECONDARY_COLOR = RGBColor(99, 102, 241)
DANGER_COLOR = RGBColor(239, 68, 68)
WARNING_COLOR = RGBColor(245, 158, 11)
DARK_COLOR = RGBColor(15, 23, 42)
WHITE_COLOR = RGBColor(255, 255, 255)

def add_content_slide(prs, title, content_items, highlight_color=PRIMARY_COLOR):
    """اضافة شريحة محتوى"""
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)
    
    # خلفية
    background = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    background.fill.solid()
    background.fill.fore_color.rgb = DARK_COLOR
    background.line.fill.background()
    
    # شريط العنوان
    title_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(1.3))
    title_bar.fill.solid()
    title_bar.fill.fore_color.rgb = highlight_color
    title_bar.line.fill.background()
    
    # العنوان
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(12.333), Inches(0.8))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = WHITE_COLOR
    p.alignment = PP_ALIGN.RIGHT
    
    # المحتوى
    y_pos = 1.8
    for item in content_items:
        content_box = slide.shapes.add_textbox(Inches(0.8), Inches(y_pos), Inches(11.733), Inches(0.8))
        tf = content_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = item
        p.font.size = Pt(24)
        p.font.color.rgb = WHITE_COLOR
        p.alignment = PP_ALIGN.RIGHT
        y_pos += 0.9
    
    return slide

# اضافة شرائح جديدة

# شريحة: البنية التقنية
add_content_slide(prs, "البنية التقنية", [
    "Python 3.14 + FastAPI",
    "LSTM Neural Network (NumPy)",
    "Digital Twin Manager",
    "Message Queue (Kafka/RabbitMQ Ready)",
    "RESTful API with Swagger Documentation",
    "Real-time Processing < 5ms",
], SECONDARY_COLOR)

# شريحة: API Endpoints
add_content_slide(prs, "API Endpoints", [
    "POST /events - تحليل حدث وارجاع تقييم المخاطر",
    "POST /events/async - معالجة غير متزامنة",
    "GET /risk/{user_id} - ملف المخاطر للمستخدم",
    "POST /feedback - تغذية راجعة للتعلم",
    "GET /stats - احصائيات النظام",
    "GET /health - فحص صحة الخادم",
], PRIMARY_COLOR)

# شريحة: عوامل تقييم المخاطر
add_content_slide(prs, "عوامل تقييم المخاطر", [
    "الموقع الجغرافي وكشف VPN/Tor (30%)",
    "بصمة الجهاز وتغييره (20%)",
    "سرعة التنقل بين الصفحات (20%)",
    "تنبؤ LSTM بالاجراء التالي (30%)",
    "قيمة الاصل (سيارة/عقار)",
    "هل المستلم معروف ام جديد",
], WARNING_COLOR)

# شريحة: خدمات ابشر المدعومة
add_content_slide(prs, "خدمات ابشر المدعومة", [
    "نقل ملكية المركبات",
    "نقل ملكية العقارات",
    "نقل ملكية السجلات التجارية",
    "تجديد الاقامات والتاشيرات",
    "اصدار جوازات السفر",
    "جميع الخدمات الحساسة",
], PRIMARY_COLOR)

# شريحة: مقاييس الاداء
add_content_slide(prs, "مقاييس الاداء المتوقعة", [
    "زمن المعالجة: < 5 مللي ثانية",
    "دقة الكشف: 99.2%",
    "نسبة الايجابيات الكاذبة: < 1%",
    "القدرة: 10,000+ حدث/ثانية",
    "التوفر: 99.99%",
    "التعلم المستمر: اسبوعيا",
], SECONDARY_COLOR)

# حفظ العرض المحدث
output_path = r'C:\Users\Administrator\Desktop\Rased_Presentation.pptx'
prs.save(output_path)
print("Done! Presentation updated successfully.")
print(f"Total slides now: {len(prs.slides)}")
