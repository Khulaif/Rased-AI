"""
إنشاء عرض بوربوينت لمشروع راصد
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from pptx.oxml import parse_xml

# إنشاء عرض جديد
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# ألوان المشروع
PRIMARY_COLOR = RGBColor(16, 185, 129)  # أخضر
SECONDARY_COLOR = RGBColor(99, 102, 241)  # بنفسجي
DANGER_COLOR = RGBColor(239, 68, 68)  # أحمر
WARNING_COLOR = RGBColor(245, 158, 11)  # أصفر
DARK_COLOR = RGBColor(15, 23, 42)  # أزرق داكن
WHITE_COLOR = RGBColor(255, 255, 255)

def add_title_slide(prs, title, subtitle):
    """إضافة شريحة العنوان الرئيسية"""
    slide_layout = prs.slide_layouts[6]  # Blank
    slide = prs.slides.add_slide(slide_layout)
    
    # خلفية
    background = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    background.fill.solid()
    background.fill.fore_color.rgb = DARK_COLOR
    background.line.fill.background()
    
    # الشعار
    logo = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(5.5), Inches(1.5), Inches(2.333), Inches(1.5))
    logo.fill.solid()
    logo.fill.fore_color.rgb = PRIMARY_COLOR
    logo.line.fill.background()
    
    # نص الشعار
    logo_text = slide.shapes.add_textbox(Inches(5.5), Inches(1.8), Inches(2.333), Inches(1))
    tf = logo_text.text_frame
    p = tf.paragraphs[0]
    p.text = "🛡️"
    p.font.size = Pt(48)
    p.alignment = PP_ALIGN.CENTER
    
    # العنوان
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(3.2), Inches(12.333), Inches(1.5))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(54)
    p.font.bold = True
    p.font.color.rgb = WHITE_COLOR
    p.alignment = PP_ALIGN.CENTER
    
    # العنوان الفرعي
    subtitle_box = slide.shapes.add_textbox(Inches(0.5), Inches(4.7), Inches(12.333), Inches(1))
    tf = subtitle_box.text_frame
    p = tf.paragraphs[0]
    p.text = subtitle
    p.font.size = Pt(24)
    p.font.color.rgb = PRIMARY_COLOR
    p.alignment = PP_ALIGN.CENTER
    
    return slide

def add_content_slide(prs, title, content_items, highlight_color=PRIMARY_COLOR):
    """إضافة شريحة محتوى"""
    slide_layout = prs.slide_layouts[6]  # Blank
    slide = prs.slides.add_slide(slide_layout)
    
    # خلفية
    background = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    background.fill.solid()
    background.fill.fore_color.rgb = DARK_COLOR
    background.line.fill.background()
    
    # شريط العنوان
    title_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(1.3))
    title_bar.fill.solid()
    title_bar.fill.fore_color.rgb = highlight_color
    title_bar.line.fill.background()
    
    # العنوان
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(12.333), Inches(0.8))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = WHITE_COLOR
    p.alignment = PP_ALIGN.RIGHT
    
    # المحتوى
    y_pos = 1.8
    for item in content_items:
        content_box = slide.shapes.add_textbox(Inches(0.8), Inches(y_pos), Inches(11.733), Inches(0.8))
        tf = content_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = f"◀  {item}"
        p.font.size = Pt(24)
        p.font.color.rgb = WHITE_COLOR
        p.alignment = PP_ALIGN.RIGHT
        y_pos += 0.9
    
    return slide

def add_scenario_slide(prs, title, scenario_title, items, risk_level, color):
    """إضافة شريحة سيناريو"""
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)
    
    # خلفية
    background = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    background.fill.solid()
    background.fill.fore_color.rgb = DARK_COLOR
    background.line.fill.background()
    
    # شريط العنوان
    title_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(1.3))
    title_bar.fill.solid()
    title_bar.fill.fore_color.rgb = color
    title_bar.line.fill.background()
    
    # العنوان
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(12.333), Inches(0.8))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = WHITE_COLOR
    p.alignment = PP_ALIGN.RIGHT
    
    # عنوان السيناريو
    scenario_box = slide.shapes.add_textbox(Inches(0.5), Inches(1.6), Inches(12.333), Inches(0.8))
    tf = scenario_box.text_frame
    p = tf.paragraphs[0]
    p.text = scenario_title
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = color
    p.alignment = PP_ALIGN.RIGHT
    
    # المحتوى
    y_pos = 2.5
    for item in items:
        content_box = slide.shapes.add_textbox(Inches(0.8), Inches(y_pos), Inches(8), Inches(0.6))
        tf = content_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = item
        p.font.size = Pt(20)
        p.font.color.rgb = WHITE_COLOR
        p.alignment = PP_ALIGN.RIGHT
        y_pos += 0.65
    
    # مربع مستوى المخاطر
    risk_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(9.5), Inches(2.5), Inches(3), Inches(3.5))
    risk_box.fill.solid()
    risk_box.fill.fore_color.rgb = color
    risk_box.line.fill.background()
    
    # نص المخاطر
    risk_text = slide.shapes.add_textbox(Inches(9.5), Inches(3), Inches(3), Inches(2.5))
    tf = risk_text.text_frame
    p = tf.paragraphs[0]
    p.text = risk_level
    p.font.size = Pt(48)
    p.font.bold = True
    p.font.color.rgb = WHITE_COLOR
    p.alignment = PP_ALIGN.CENTER
    
    return slide

# ===============================
# إنشاء الشرائح
# ===============================

# 1. شريحة العنوان
add_title_slide(
    prs,
    "راصد - نظام الكشف عن الاحتيال",
    "حماية عمليات نقل الملكية في منصة أبشر بالذكاء الاصطناعي"
)

# 2. المشكلة
add_content_slide(prs, "المشكلة", [
    "محاولات سرقة ملكية السيارات والعقارات عبر منصة أبشر",
    "استخدام بيانات مسروقة لنقل الملكيات بدون علم المالك",
    "هجمات البوتات الآلية بسرعة تنقل غير بشرية",
    "صعوبة التمييز بين المستخدم الحقيقي والمحتال",
    "خسائر مالية ضخمة للمواطنين والمقيمين",
], DANGER_COLOR)

# 3. الحل
add_content_slide(prs, "الحل: نظام راصد", [
    "نظام ذكاء اصطناعي يعمل في الوقت الفعلي",
    "يحلل سلوك المستخدم ويقارنه بنمطه المعتاد",
    "يكتشف الشذوذ في الموقع، الجهاز، وسرعة التصفح",
    "يتخذ قرارات فورية: السماح / التحقق / الحظر",
    "يتعلم باستمرار من التغذية الراجعة",
], PRIMARY_COLOR)

# 4. كيف يعمل راصد
add_content_slide(prs, "كيف يعمل راصد؟", [
    "🧠  شبكة LSTM العصبية تتنبأ بالإجراء التالي للمستخدم",
    "👤  التوأم الرقمي يحفظ البصمة السلوكية لكل مستخدم",
    "📍  تحليل الموقع الجغرافي وكشف VPN/Tor",
    "📱  بصمة الجهاز والتحقق من تغييره",
    "⚡  سرعة المعالجة أقل من 5 مللي ثانية",
], SECONDARY_COLOR)

# 5. سيناريو 1 - أخضر
add_scenario_slide(
    prs,
    "السيناريو 1: مستخدم طبيعي",
    "✅ تصفح المركبات من الجهاز المعتاد",
    [
        "المستخدم: user_00000",
        "الإجراء: تصفح المركبات",
        "الموقع: الخُبر 🇸🇦",
        "الجهاز: iPhone (iOS)",
        "",
        "القرار: ✅ مسموح",
        "بدون أي إزعاج للمستخدم",
    ],
    "25/100",
    PRIMARY_COLOR
)

# 6. سيناريو 2 - أصفر
add_scenario_slide(
    prs,
    "السيناريو 2: نقل ملكية مشبوه",
    "⚠️ محاولة نقل ملكية سيارة من موقع غير معتاد",
    [
        "نفس المستخدم يظهر فجأة من روسيا 🇷🇺",
        "جهاز مختلف (Windows بدل iPhone)",
        "استخدام VPN مكتشف",
        "محاولة نقل ملكية سيارة بقيمة 85,000 ريال",
        "",
        "القرار: ⚠️ مطلوب تحقق إضافي",
        "أسئلة أمان / OTP / نفاذ",
    ],
    "52/100",
    WARNING_COLOR
)

# 7. سيناريو 3 - أحمر
add_scenario_slide(
    prs,
    "السيناريو 3: هجوم بوت - سرقة عقار",
    "🛑 محاولة سرقة عقار بقيمة 1.5 مليون ريال",
    [
        "موقع: لاغوس، نيجيريا 🇳🇬",
        "استخدام Tor (شبكة مظلمة)",
        "سرعة تنقل: 0.02 ثانية (غير بشرية!)",
        "محاولة تأكيد نقل ملكية عقار",
        "",
        "القرار: 🚫 محظور فوراً",
        "تنبيه SOC + إشعار المالك الحقيقي",
    ],
    "86/100",
    DANGER_COLOR
)

# 8. التعلم المستمر
add_content_slide(prs, "التعلم المستمر", [
    "🔄  إذا اجتاز المستخدم التحقق ← النظام يتعلم أنه سلوك صحيح",
    "📉  يخفض نقاط المخاطر المستقبلية لهذا النمط",
    "🎯  إذا تأكد الاحتيال ← يزيد الحساسية لهذا النمط",
    "🧠  إعادة تدريب النموذج أسبوعياً",
    "📊  تقارير أداء لتحسين الدقة باستمرار",
], SECONDARY_COLOR)

# 9. المميزات
add_content_slide(prs, "مميزات نظام راصد", [
    "⚡  سرعة فائقة: أقل من 5 مللي ثانية",
    "🎯  دقة عالية: 99.2% في كشف الاحتيال",
    "👻  غير مرئي: لا يؤثر على تجربة المستخدم الطبيعي",
    "🔒  حماية شاملة: سيارات، عقارات، سجلات تجارية",
    "📱  متوافق مع جميع الأجهزة والمتصفحات",
], PRIMARY_COLOR)

# 10. الخلاصة
add_title_slide(
    prs,
    "راصد - نحمي ملكياتك الرقمية",
    "شكراً لاستماعكم 🛡️"
)

# حفظ العرض
output_path = r"C:\Users\Administrator\Desktop\Rased\Rased_Presentation.pptx"
prs.save(output_path)
print(f"تم إنشاء العرض التقديمي: {output_path}")
